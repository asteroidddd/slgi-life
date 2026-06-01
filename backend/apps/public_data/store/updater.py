"""Store public-data updater for the current capston schema."""

from __future__ import annotations

import hashlib
import json
import os
import re
import socket
import time as sleep_time
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import Request, urlopen

from django.contrib.gis.geos import Point
from django.db import transaction
from openpyxl import load_workbook

from apps.public_data.api_keys import env_key_ring, with_rate_limit_fallback
from apps.public_data.exceptions import RateLimitedError, is_rate_limited_code, is_rate_limited_text
from apps.public_data.regions.models import Adong, Ldong
from apps.public_data.state import dataset_state, load_state, record_dataset_result, save_state
from apps.public_data.store.models import BusinessCategory, DaisoStore, KsciCategory, Store
from apps.common.geocoding import geocode_address


DATA_DIR = Path(__file__).resolve().parents[3] / "data"
BUSINESS_CATEGORY_PATH = DATA_DIR / "store_business_category.xlsx"
KSCI_CATEGORY_PATH = DATA_DIR / "KSIC_10th.xlsx"
STORE_API_URL = "https://apis.data.go.kr/B553077/api/open/sdsc2/storeListInDong"
DAISO_BASE_URL = "https://www.daiso.co.kr"
DAISO_SEOUL = "서울"
PAGE_SIZE = 1000
MEDICAL_STORE_CATEGORY_CODES = {"G21501"}
MEDICAL_STORE_MAIN_CATEGORY_CODES = {"Q1"}
DAISO_REQUEST_DELAY_SECONDS = 0.15
DAISO_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/125.0 Safari/537.36"
)


@dataclass(frozen=True)
class StoresUpdateOptions:
    dry_run: bool = True
    force: bool = False
    limit: int | None = None
    delete_missing: bool = True
    request_interval_seconds: float = 0.2
    request_timeout_seconds: float = 40.0


@dataclass(frozen=True)
class DaisoCrawlStore:
    name: str
    address: str
    latitude: str = ""
    longitude: str = ""


class DaisoStoreHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.stores: list[DaisoCrawlStore] = []
        self._in_store = False
        self._store_depth = 0
        self._current_attrs: dict[str, str] = {}
        self._field: str | None = None
        self._name_parts: list[str] = []
        self._address_parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attrs_dict = {key: value or "" for key, value in attrs}
        class_names = set(attrs_dict.get("class", "").split())

        if tag == "div" and "bx-store" in class_names:
            self._in_store = True
            self._store_depth = 1
            self._current_attrs = attrs_dict
            self._field = None
            self._name_parts = []
            self._address_parts = []
            return

        if not self._in_store:
            return

        if tag == "div":
            self._store_depth += 1
        if tag == "h4" and "place" in class_names:
            self._field = "name"
        elif tag == "p" and "addr" in class_names:
            self._field = "address"

    def handle_endtag(self, tag: str) -> None:
        if not self._in_store:
            return
        if tag in {"h4", "p"}:
            self._field = None
        if tag == "div":
            self._store_depth -= 1
            if self._store_depth == 0:
                self._finish_store()

    def handle_data(self, data: str) -> None:
        if not self._in_store or self._field is None:
            return
        cleaned = _normalize_text(data)
        if not cleaned:
            return
        if self._field == "name":
            self._name_parts.append(cleaned)
        elif self._field == "address":
            self._address_parts.append(cleaned)

    def _finish_store(self) -> None:
        name = _normalize_text(" ".join(self._name_parts))
        address = _normalize_text(" ".join(self._address_parts))
        if name and address:
            self.stores.append(
                DaisoCrawlStore(
                    name=name,
                    address=address,
                    latitude=self._current_attrs.get("data-lat", ""),
                    longitude=self._current_attrs.get("data-lng", ""),
                )
            )
        self._in_store = False
        self._store_depth = 0
        self._current_attrs = {}
        self._field = None


def _require_env(name: str) -> str:
    value = os.environ.get(name, "").strip()
    if not value:
        raise RuntimeError(f"{name} is required")
    return value


def _file_meta(path: Path) -> dict[str, Any]:
    data = path.read_bytes()
    return {
        "file": str(path),
        "file_hash": hashlib.sha256(data).hexdigest(),
        "file_size": len(data),
    }


def _catalog_snapshot() -> dict[str, Any]:
    business = _file_meta(BUSINESS_CATEGORY_PATH)
    ksci = _file_meta(KSCI_CATEGORY_PATH)
    digest = hashlib.sha256()
    for item in (business, ksci):
        digest.update(item["file"].encode("utf-8"))
        digest.update(b"\0")
        digest.update(item["file_hash"].encode("ascii"))
        digest.update(b"\0")
    return {
        "file_hash": digest.hexdigest(),
        "files": {
            "store_business_category.xlsx": business,
            "KSIC_10th.xlsx": ksci,
        },
    }


def _business_rows() -> list[dict[str, str]]:
    wb = load_workbook(BUSINESS_CATEGORY_PATH, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        main_code, main_name, middle_code, middle_name, sub_code, sub_name = row[:6]
        if not sub_code:
            continue
        rows.append(
            {
                "main_code": str(main_code).strip(),
                "main_name": str(main_name or "").strip(),
                "middle_code": str(middle_code).strip(),
                "middle_name": str(middle_name or "").strip(),
                "sub_code": str(sub_code).strip(),
                "sub_name": str(sub_name or "").strip(),
            }
        )
    if len(rows) < 200:
        raise RuntimeError(f"refusing incomplete business category file: {len(rows)} rows")
    return rows


def _ksci_rows() -> list[dict[str, str]]:
    wb = load_workbook(KSCI_CATEGORY_PATH, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    current = {"main_code": "", "main_name": "", "middle_name": "", "sub_name": "", "class_name": ""}
    rows = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        main_code, main_name, _middle_code, middle_name, _sub_code, sub_name, _class_code, class_name, detail_code, detail_name = row[:10]
        if main_code:
            current["main_code"] = str(main_code).strip()
            current["main_name"] = str(main_name or "").strip()
        if middle_name:
            current["middle_name"] = str(middle_name).strip()
        if sub_name:
            current["sub_name"] = str(sub_name).strip()
        if class_name:
            current["class_name"] = str(class_name).strip()
        if not detail_code:
            continue
        rows.append(
            {
                "ksci_code": f"{current['main_code']}{str(detail_code).strip()}",
                "subcategory_name": str(detail_name or "").strip(),
                "class_name": current["class_name"],
                "subclass_name": current["sub_name"],
                "middle_category_name": current["middle_name"],
                "main_category_name": current["main_name"],
            }
        )
    if len(rows) < 1000:
        raise RuntimeError(f"refusing incomplete KSIC file: {len(rows)} rows")
    return rows


def _catalog_tables_populated() -> bool:
    return BusinessCategory.objects.exists() and KsciCategory.objects.exists()


def _load_catalog(options: StoresUpdateOptions, snapshot: dict[str, Any]) -> dict[str, Any]:
    previous_hash = dataset_state("stores").get("catalog_snapshot", {}).get("file_hash")
    skipped_write = bool(
        previous_hash == snapshot["file_hash"]
        and not options.force
        and not options.dry_run
        and _catalog_tables_populated()
    )
    if skipped_write:
        return {
            "status": "success",
            "completed": True,
            "checked": {"business_category": 0, "ksci_category": 0},
            "loaded": {"business_category": 0, "ksci_category": 0},
            "skipped_write": True,
            "reason": "file_hash_unchanged",
            "files": snapshot,
        }

    business_rows = _business_rows()
    ksci_rows = _ksci_rows()
    loaded_business = loaded_ksci = 0

    if not options.dry_run:
        with transaction.atomic():
            BusinessCategory.objects.bulk_create(
                [
                    BusinessCategory(
                        subcategory_code=row["sub_code"],
                        subcategory_name=row["sub_name"],
                        middle_category_code=row["middle_code"],
                        middle_category_name=row["middle_name"],
                        main_category_code=row["main_code"],
                        main_category_name=row["main_name"],
                    )
                    for row in business_rows
                ],
                batch_size=1000,
                update_conflicts=True,
                update_fields=[
                    "subcategory_name",
                    "middle_category_code",
                    "middle_category_name",
                    "main_category_code",
                    "main_category_name",
                ],
                unique_fields=["subcategory_code"],
            )
            loaded_business = len(business_rows)

            KsciCategory.objects.bulk_create(
                [
                    KsciCategory(
                        ksci_code=row["ksci_code"],
                        subcategory_name=row["subcategory_name"],
                        class_name=row["class_name"],
                        subclass_name=row["subclass_name"],
                        middle_category_name=row["middle_category_name"],
                        main_category_name=row["main_category_name"],
                    )
                    for row in ksci_rows
                ],
                batch_size=1000,
                update_conflicts=True,
                update_fields=[
                    "subcategory_name",
                    "class_name",
                    "subclass_name",
                    "middle_category_name",
                    "main_category_name",
                ],
                unique_fields=["ksci_code"],
            )
            loaded_ksci = len(ksci_rows)

    return {
        "status": "success",
        "completed": True,
        "checked": {"business_category": len(business_rows), "ksci_category": len(ksci_rows)},
        "loaded": {"business_category": loaded_business, "ksci_category": loaded_ksci},
        "skipped_write": False,
        "files": snapshot,
    }


def _request_json(params: dict[str, str], options: StoresUpdateOptions) -> dict[str, Any]:
    if options.request_interval_seconds:
        sleep_time.sleep(options.request_interval_seconds)
    req = Request(
        f"{STORE_API_URL}?{urlencode(params, safe='%')}",
        headers={"User-Agent": "capston-public-data-updater/0.1"},
    )
    try:
        with urlopen(req, timeout=options.request_timeout_seconds) as res:
            return json.loads(res.read().decode("utf-8", errors="replace"))
    except HTTPError as exc:
        if exc.code == 429:
            raise RateLimitedError("store API rate limited: HTTP 429") from exc
        raise RuntimeError(f"store API request failed: HTTP {exc.code}") from exc
    except (URLError, TimeoutError, socket.timeout) as exc:
        raise RuntimeError(f"store API request failed: {type(exc).__name__}") from exc
    except json.JSONDecodeError as exc:
        raise RuntimeError("store API returned malformed JSON") from exc


def _response_items(payload: dict[str, Any]) -> tuple[int, list[dict[str, Any]]]:
    header = payload.get("header") or {}
    result_code = str(header.get("resultCode") or "").strip()
    if result_code and result_code != "00":
        if is_rate_limited_code(result_code) or is_rate_limited_text(header):
            raise RateLimitedError(f"store API rate limited resultCode={result_code} resultMsg={header.get('resultMsg')}")
        raise RuntimeError(f"store API error resultCode={result_code} resultMsg={header.get('resultMsg')}")
    body = payload.get("body") or {}
    total = int(body.get("totalCount") or 0)
    items = body.get("items") or []
    if isinstance(items, dict):
        items = items.get("item") or []
    if isinstance(items, dict):
        items = [items]
    return total, items if isinstance(items, list) else []


def _fetch_gu_page(
    api_key: str,
    gu_code: str,
    page: int,
    options: StoresUpdateOptions,
) -> tuple[int, list[dict[str, Any]]]:
    payload = _request_json(
        {
            "serviceKey": api_key,
            "divId": "signguCd",
            "key": gu_code,
            "pageNo": str(page),
            "numOfRows": str(PAGE_SIZE),
            "type": "json",
        },
        options,
    )
    return _response_items(payload)


def _fetch_gu_rows(api_key: str, gu_code: str, options: StoresUpdateOptions) -> list[dict[str, Any]]:
    page = 1
    collected: list[dict[str, Any]] = []
    while True:
        total, rows = _fetch_gu_page(api_key, gu_code, page, options)
        collected.extend(rows)
        if not rows or not total or page * PAGE_SIZE >= total:
            return collected
        page += 1


def _fetch_gu_rows_with_fallback(api_keys: tuple[str, ...], gu_code: str, options: StoresUpdateOptions) -> list[dict[str, Any]]:
    return with_rate_limit_fallback(api_keys, lambda api_key: _fetch_gu_rows(api_key, gu_code, options))


def _iter_gu_rows_with_fallback(api_keys: tuple[str, ...], gu_code: str, options: StoresUpdateOptions):
    page = 1
    while True:
        total, rows = with_rate_limit_fallback(
            api_keys,
            lambda api_key: _fetch_gu_page(api_key, gu_code, page, options),
        )
        for row in rows:
            yield row
        if not rows or not total or page * PAGE_SIZE >= total:
            return
        page += 1


def _normalize_adong_code(value: Any) -> str | None:
    text = str(value or "").strip()
    if not text:
        return None
    if len(text) == 8:
        return f"{text}00"
    return text


def _point_from_values(lon: Any, lat: Any) -> Point | None:
    try:
        return Point(float(lon), float(lat), srid=4326)
    except (TypeError, ValueError):
        return None


def _point(row: dict[str, Any]) -> Point | None:
    return _point_from_values(
        row.get("lon") or row.get("lonVal") or row.get("x"),
        row.get("lat") or row.get("latVal") or row.get("y"),
    )


def _is_medical_store_category(category_id: str, category_main_category_code: str) -> bool:
    return (
        category_id in MEDICAL_STORE_CATEGORY_CODES
        or category_id.startswith("Q1")
        or category_main_category_code in MEDICAL_STORE_MAIN_CATEGORY_CODES
    )


def _find_region(point: Point, ldong_id: str | None, adong_id: str | None) -> tuple[Ldong | None, Adong | None]:
    ldong = Ldong.objects.filter(ldong_code=ldong_id).first() if ldong_id else None
    adong = Adong.objects.filter(adong_code=adong_id).first() if adong_id else None
    if not ldong:
        ldong = Ldong.objects.filter(boundary__covers=point).order_by("area_m2").first()
    if not adong:
        adong = Adong.objects.filter(boundary__covers=point).order_by("area_m2").first()
    return ldong, adong


def _build_store_record(
    row: dict[str, Any],
    *,
    category_ids: set[str],
    ksci_ids: set[str],
) -> tuple[dict[str, Any] | None, str | None, str | None]:
    store_id = str(row.get("bizesId") or "").strip()
    if not store_id:
        return None, None, "missing_id"
    point = _point(row)
    if not point:
        return None, store_id, "missing_location"
    category_id = str(row.get("indsSclsCd") or "").strip()
    category_main_category_code = str(row.get("indsLclsCd") or "").strip()
    if _is_medical_store_category(category_id, category_main_category_code):
        return None, store_id, "medical_category"
    if category_id not in category_ids:
        return None, store_id, "unknown_category"
    ldong_id = str(row.get("ldongCd") or "").strip() or None
    adong_id = _normalize_adong_code(row.get("adongCd"))
    ldong, adong = _find_region(point, ldong_id, adong_id)
    if not (ldong and adong):
        return None, store_id, "region_not_found"
    ksci_id = str(row.get("ksicCd") or "").strip() or None
    if ksci_id and ksci_id not in ksci_ids:
        ksci_id = None

    return (
        {
            "id": store_id[:50],
            "defaults": {
                "name": str(row.get("bizesNm") or "")[:100],
                "branch_name": str(row.get("brchNm") or "")[:100] or None,
                "address": str(row.get("rdnmAdr") or row.get("lnoAdr") or "")[:255],
                "location": point,
                "category_id": category_id,
                "ksci_id": ksci_id,
                "ldong_id": ldong.ldong_code,
                "adong_id": adong.adong_code,
            },
        },
        store_id[:50],
        None,
    )


def _fetch_and_build_stores(options: StoresUpdateOptions) -> dict[str, Any]:
    api_keys = env_key_ring("PUBLIC_DATA_API_KEY")
    gu_codes = list(Ldong.objects.values_list("gu_id", flat=True).distinct().order_by("gu_id"))
    category_ids = set(BusinessCategory.objects.values_list("subcategory_code", flat=True))
    ksci_ids = set(KsciCategory.objects.values_list("ksci_code", flat=True))
    if not gu_codes:
        raise RuntimeError("regions must be loaded before updating stores")
    if not category_ids:
        raise RuntimeError("business categories must be loaded before updating stores")

    checked = skipped = 0
    source_ids: set[str] = set()
    records: list[dict[str, Any]] = []
    skip_reasons = {
        "missing_id": 0,
        "missing_location": 0,
        "unknown_category": 0,
        "region_not_found": 0,
        "medical_category": 0,
    }
    completed = False

    for gu_code in gu_codes:
        rows = _fetch_gu_rows_with_fallback(api_keys, gu_code, options)
        for row in rows:
            if options.limit is not None and checked >= options.limit:
                return {
                    "status": "partial",
                    "completed": False,
                    "checked": checked,
                    "records": records,
                    "source_ids": source_ids,
                    "skipped": skipped,
                    "skip_reasons": skip_reasons,
                    "gu_count": len(gu_codes),
                }
            checked += 1
            record, source_id, reason = _build_store_record(row, category_ids=category_ids, ksci_ids=ksci_ids)
            if source_id:
                source_ids.add(source_id)
            if not record:
                skipped += 1
                if reason:
                    skip_reasons[reason] += 1
                continue
            records.append(record)
    completed = True
    return {
        "status": "success",
        "completed": completed,
        "checked": checked,
        "records": records,
        "source_ids": source_ids,
        "skipped": skipped,
        "skip_reasons": skip_reasons,
        "gu_count": len(gu_codes),
        "public_data_key_count": len(api_keys),
    }


def _upsert_store_records(records: list[dict[str, Any]]) -> dict[str, int]:
    records_by_id = {record["id"]: record for record in records}
    records = list(records_by_id.values())
    record_ids = list(records_by_id)
    if not records:
        return {"loaded": 0, "created": 0, "updated": 0}

    existing_ids = set(Store.objects.filter(id__in=record_ids).values_list("id", flat=True))
    Store.objects.bulk_create(
        [Store(id=record["id"], **record["defaults"]) for record in records],
        batch_size=2000,
        update_conflicts=True,
        update_fields=[
            "name",
            "branch_name",
            "address",
            "location",
            "category",
            "ksci",
            "ldong",
            "adong",
        ],
        unique_fields=["id"],
    )
    return {
        "loaded": len(records),
        "created": len(set(record_ids) - existing_ids),
        "updated": len(set(record_ids) & existing_ids),
    }


def _stream_and_write_stores(options: StoresUpdateOptions) -> dict[str, Any]:
    api_keys = env_key_ring("PUBLIC_DATA_API_KEY")
    gu_codes = list(Ldong.objects.values_list("gu_id", flat=True).distinct().order_by("gu_id"))
    category_ids = set(BusinessCategory.objects.values_list("subcategory_code", flat=True))
    ksci_ids = set(KsciCategory.objects.values_list("ksci_code", flat=True))
    if not gu_codes:
        raise RuntimeError("regions must be loaded before updating stores")
    if not category_ids:
        raise RuntimeError("business categories must be loaded before updating stores")

    checked = loaded = created = updated = skipped = deleted_missing = 0
    source_ids: set[str] = set()
    seen_ids: set[str] = set()
    batch: list[dict[str, Any]] = []
    skip_reasons = {
        "missing_id": 0,
        "missing_location": 0,
        "unknown_category": 0,
        "region_not_found": 0,
        "medical_category": 0,
    }

    def flush_batch() -> None:
        nonlocal loaded, created, updated, batch
        if not batch:
            return
        if options.dry_run:
            loaded += len(batch)
        else:
            with transaction.atomic():
                stats = _upsert_store_records(batch)
            loaded += stats["loaded"]
            created += stats["created"]
            updated += stats["updated"]
        batch = []

    for gu_code in gu_codes:
        for row in _iter_gu_rows_with_fallback(api_keys, gu_code, options):
            if options.limit is not None and checked >= options.limit:
                flush_batch()
                return {
                    "status": "partial",
                    "completed": False,
                    "checked": checked,
                    "loaded": loaded,
                    "created": created,
                    "updated": updated,
                    "skipped": skipped,
                    "deleted_missing": 0,
                    "dry_run": options.dry_run,
                    "gu_count": len(gu_codes),
                    "skip_reasons": skip_reasons,
                    "delete_missing": False,
                    "reason": "limit_reached",
                }
            checked += 1
            record, source_id, reason = _build_store_record(row, category_ids=category_ids, ksci_ids=ksci_ids)
            if options.delete_missing and source_id:
                source_ids.add(source_id)
            if not record:
                skipped += 1
                if reason:
                    skip_reasons[reason] += 1
                continue
            if record["id"] in seen_ids:
                continue
            seen_ids.add(record["id"])
            batch.append(record)
            if len(batch) >= 2000:
                flush_batch()

    flush_batch()
    if not options.dry_run and options.delete_missing:
        missing_qs = Store.objects.exclude(id__in=source_ids)
        deleted_missing = missing_qs.count()
        missing_qs.delete()

    return {
        "status": "success",
        "completed": True,
        "checked": checked,
        "loaded": loaded,
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "deleted_missing": deleted_missing,
        "dry_run": options.dry_run,
        "gu_count": len(gu_codes),
        "skip_reasons": skip_reasons,
        "delete_missing": options.delete_missing,
    }


def _normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def _daiso_request_text(path: str, params: dict[str, str], options: StoresUpdateOptions) -> str:
    if DAISO_REQUEST_DELAY_SECONDS:
        sleep_time.sleep(DAISO_REQUEST_DELAY_SECONDS)
    req = Request(
        f"{DAISO_BASE_URL}{path}?{urlencode(params)}",
        headers={"User-Agent": DAISO_USER_AGENT},
    )
    try:
        with urlopen(req, timeout=options.request_timeout_seconds) as response:
            charset = response.headers.get_content_charset() or "utf-8"
            return response.read().decode(charset, errors="replace")
    except (HTTPError, URLError, TimeoutError, socket.timeout) as exc:
        raise RuntimeError(f"Daiso request failed: {type(exc).__name__}") from exc


def _daiso_request_json(path: str, params: dict[str, str], options: StoresUpdateOptions) -> list[dict[str, Any]]:
    body = _daiso_request_text(path, params, options)
    try:
        parsed = json.loads(body)
    except json.JSONDecodeError as exc:
        raise RuntimeError("Daiso returned malformed JSON") from exc
    if not isinstance(parsed, list):
        raise RuntimeError(f"Daiso returned unexpected JSON: {type(parsed).__name__}")
    return parsed


def _daiso_get_values(path: str, params: dict[str, str], options: StoresUpdateOptions) -> list[str]:
    rows = _daiso_request_json(path, params, options)
    values = [_normalize_text(str(row.get("value", ""))) for row in rows]
    return [value for value in values if value]


def _parse_daiso_stores(html: str) -> list[DaisoCrawlStore]:
    parser = DaisoStoreHTMLParser()
    parser.feed(html)
    parser.close()
    return parser.stores


def _crawl_daiso_stores(options: StoresUpdateOptions) -> list[DaisoCrawlStore]:
    stores_by_key: dict[tuple[str, str], DaisoCrawlStore] = {}
    districts = _daiso_get_values("/cs/ajax/sido_search", {"sido": DAISO_SEOUL}, options)
    for district in districts:
        subdistricts = _daiso_get_values(
            "/cs/ajax/gugun_search",
            {"sido": DAISO_SEOUL, "gugun": district},
            options,
        )
        if not subdistricts:
            subdistricts = [""]
        for subdistrict in subdistricts:
            params = {"sido": DAISO_SEOUL, "gugun": district}
            if subdistrict:
                params["dong"] = subdistrict
            html = _daiso_request_text("/cs/ajax/shop_search", params, options)
            for store in _parse_daiso_stores(html):
                stores_by_key[(store.name, store.address)] = store
    return sorted(stores_by_key.values(), key=lambda item: (item.name, item.address))


def _daiso_id(name: str, address: str) -> str:
    digest = hashlib.sha1(f"{name}\0{address}".encode("utf-8")).hexdigest()
    return f"daiso_{digest[:32]}"


def _daiso_location(store: DaisoCrawlStore, options: StoresUpdateOptions) -> tuple[Point | None, str]:
    point = _point_from_values(store.longitude, store.latitude)
    if point:
        return point, "source_coord"
    result = geocode_address(
        store.address,
        request_timeout=min(options.request_timeout_seconds, 10.0),
        request_interval=options.request_interval_seconds,
        user_agent="capston-daiso-updater/0.1",
        prefer_kakao=True,
    )
    if result.status == "success" and result.point:
        return result.point, result.provider or "geocode"
    return None, result.error or "geocode_failed"


def _upsert_daiso_records(records: list[dict[str, Any]]) -> dict[str, int]:
    records_by_id = {record["id"]: record for record in records}
    records = list(records_by_id.values())
    record_ids = list(records_by_id)
    if not records:
        return {"loaded": 0, "created": 0, "updated": 0}
    existing_ids = set(DaisoStore.objects.filter(id__in=record_ids).values_list("id", flat=True))
    DaisoStore.objects.bulk_create(
        [DaisoStore(id=record["id"], **record["defaults"]) for record in records],
        batch_size=500,
        update_conflicts=True,
        update_fields=["name", "address", "location"],
        unique_fields=["id"],
    )
    return {
        "loaded": len(records),
        "created": len(set(record_ids) - existing_ids),
        "updated": len(set(record_ids) & existing_ids),
    }


def update_daiso_stores(options: StoresUpdateOptions) -> dict[str, Any]:
    crawled = _crawl_daiso_stores(options)
    if not crawled:
        raise RuntimeError("refusing empty Daiso crawl result")
    checked = loaded = created = updated = skipped = deleted_missing = 0
    source_ids: set[str] = set()
    records: list[dict[str, Any]] = []
    skip_reasons: dict[str, int] = {
        "missing_location": 0,
    }
    location_sources: dict[str, int] = {}

    for store in crawled:
        checked += 1
        store_id = _daiso_id(store.name, store.address)
        source_ids.add(store_id)
        point, location_source = _daiso_location(store, options)
        location_sources[location_source] = location_sources.get(location_source, 0) + 1
        if not point:
            skipped += 1
            skip_reasons["missing_location"] += 1
            continue
        records.append(
            {
                "id": store_id,
                "defaults": {
                    "name": store.name[:200],
                    "address": store.address[:255],
                    "location": point,
                },
            }
        )
        if options.limit is not None and checked >= options.limit:
            break

    if options.dry_run:
        loaded = len(records)
    else:
        with transaction.atomic():
            stats = _upsert_daiso_records(records)
            loaded = stats["loaded"]
            created = stats["created"]
            updated = stats["updated"]
        if options.delete_missing:
            missing_qs = DaisoStore.objects.exclude(id__in=source_ids)
            deleted_missing = missing_qs.count()
            missing_qs.delete()

    result = {
        "dataset": "daiso",
        "dry_run": options.dry_run,
        "status": "success",
        "completed": options.limit is None or checked < options.limit,
        "checked": checked,
        "loaded": loaded,
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "deleted_missing": deleted_missing,
        "skip_reasons": skip_reasons,
        "location_sources": location_sources,
    }
    if options.limit is not None and checked >= options.limit:
        result["status"] = "partial"
        result["reason"] = "limit_reached"
    if not options.dry_run:
        record_dataset_result("daiso", result)
    return result


def update_stores_data(options: StoresUpdateOptions) -> dict[str, Any]:
    return _stream_and_write_stores(options)


def update_stores(options: StoresUpdateOptions) -> dict[str, Any]:
    catalog_snapshot = _catalog_snapshot()
    catalog = _load_catalog(options, catalog_snapshot)
    stores = update_stores_data(options)
    completed = bool(catalog["completed"] and stores["completed"])
    catalog_loaded = sum(catalog["loaded"].values())
    return {
        "status": "success" if completed else "partial",
        "completed": completed,
        "loaded": catalog_loaded + stores["loaded"],
        "dry_run": options.dry_run,
        "catalog": catalog,
        "stores": stores,
    }


def update(options: StoresUpdateOptions) -> dict[str, Any]:
    result: dict[str, Any] = {
        "dataset": "stores",
        "dry_run": options.dry_run,
        "stores": None,
    }
    try:
        result["stores"] = update_stores(options)
    except Exception as exc:
        result["error"] = {
            "type": type(exc).__name__,
            "message": str(exc),
        }
        if not options.dry_run:
            record_dataset_result("stores", result)
        raise

    if not options.dry_run:
        record_dataset_result("stores", result)
        if result["stores"] and result["stores"].get("completed"):
            state = load_state()
            stores_state = state.setdefault("datasets", {}).setdefault("stores", {})
            if not result["stores"]["catalog"].get("skipped_write"):
                stores_state["catalog_snapshot"] = result["stores"]["catalog"]["files"] | {
                    "loaded": result["stores"]["catalog"]["loaded"],
                }
            stores_state["stores_snapshot"] = {
                "last_success_date": None,
                "checked": result["stores"]["stores"]["checked"],
                "loaded": result["stores"]["stores"]["loaded"],
                "deleted_missing": result["stores"]["stores"]["deleted_missing"],
                "skip_reasons": result["stores"]["stores"]["skip_reasons"],
            }
            save_state(state)
    return result
